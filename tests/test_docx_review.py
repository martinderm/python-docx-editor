"""Behavioral contract tests for the additive ``docx-review.v1`` output.

The tests exercise only the public CLI seams named in the confirmed specification:

* ``python scripts/docx_ops.py inspect --in <docx> --json``
* ``python scripts/extract_docx_for_llm.py --in <docx> --out <json> [--rag-output <json>]``

Every fixture is a tiny DOCX package generated at runtime with python-docx and
direct ZIP/XML edits.  Expected values (text views, codepoint offsets, XML
locations, metadata) are literals derived from the specification, never from
the product's own extraction code.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
import textwrap
import unittest
import zipfile
from pathlib import Path

from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[1]
INSPECT_CLI = REPO_ROOT / "scripts" / "docx_ops.py"
EXTRACT_CLI = REPO_ROOT / "scripts" / "extract_docx_for_llm.py"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
W15_NS = "http://schemas.microsoft.com/office/word/2012/wordml"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
V_NS = "urn:schemas-microsoft-com:vml"

DOCUMENT_PART = "word/document.xml"
REVIEW_SCHEMA = "docx-review.v1"


def _location(path: str) -> dict:
    return {"part": DOCUMENT_PART, "path": path}


P1 = "/document[1]/body[1]/p[1]"
P2 = "/document[1]/body[1]/p[2]"
P3 = "/document[1]/body[1]/p[3]"
P4 = "/document[1]/body[1]/p[4]"
TABLE_CELL_P = "/document[1]/body[1]/tbl[1]/tr[1]/tc[1]/p[1]"
NESTED_TABLE_CELL_P = (
    "/document[1]/body[1]/tbl[1]/tr[1]/tc[1]/tbl[1]/tr[1]/tc[1]/p[1]"
)

# ---------------------------------------------------------------------------
# Shared fixture bodies
# ---------------------------------------------------------------------------

MIXED_BODY = textwrap.dedent(
    """\
    <w:p><w:r><w:t>Intro</w:t></w:r></w:p>
    <w:p><w:r><w:t>A</w:t></w:r><w:commentRangeStart w:id="1"/>
    <w:ins w:id="10" w:author="Alice" w:date="2026-10-01T10:00:00Z"><w:r><w:t>B</w:t></w:r><w:r><w:t>C</w:t></w:r></w:ins>
    <w:commentRangeEnd w:id="1"/><w:r><w:commentReference w:id="1"/></w:r>
    <w:r><w:t>D</w:t></w:r>
    <w:del w:id="11" w:author="Bob" w:date="2026-10-02T11:30:00Z"><w:r><w:delText>E</w:delText></w:r><w:r><w:delText>F</w:delText></w:r></w:del>
    <w:r><w:t>G</w:t></w:r></w:p>
    <w:p/>
    <w:p><w:ins w:id="13" w:author="Carol"><w:r><w:t>Nur</w:t></w:r></w:ins></w:p>
    """
)

MIXED_COMMENTS = textwrap.dedent(
    """\
    <w:comment w:id="1" w:author="Alice" w:date="2026-10-01T10:00:00Z">
      <w:p w14:paraId="AAAA1111"><w:r><w:t>Parent note</w:t></w:r></w:p>
    </w:comment>
    """
)

MIXED_COMMENTS_EXTENDED = (
    '<w15:commentEx w15:paraId="AAAA1111" w15:done="1"/>'
)

MIXED_PARAGRAPH_2 = {
    "location": _location(P2),
    "text_original": "ADEFG",
    "text_proposed": "ABCDG",
    "text_markup": "ABCDEFG",
}

MIXED_REVISION_INSERT = {
    "kind": "insert",
    "id": "10",
    "author": "Alice",
    "date": "2026-10-01T10:00:00Z",
    "text": "BC",
    "location": _location(P2),
    "ranges": {
        "original": [1, 1],
        "proposed": [1, 3],
        "markup": [1, 3],
    },
    "missing_metadata": [],
}

MIXED_REVISION_DELETE = {
    "kind": "delete",
    "id": "11",
    "author": "Bob",
    "date": "2026-10-02T11:30:00Z",
    "text": "EF",
    "location": _location(P2),
    "ranges": {
        "original": [2, 4],
        "proposed": [4, 4],
        "markup": [4, 6],
    },
    "missing_metadata": [],
}

MIXED_REVISION_DATE_MISSING = {
    "kind": "insert",
    "id": "13",
    "author": "Carol",
    "date": None,
    "text": "Nur",
    "location": _location(P4),
    "ranges": {
        "original": [0, 0],
        "proposed": [0, 3],
        "markup": [0, 3],
    },
    "missing_metadata": ["date"],
}


class DocxReviewTestBase(unittest.TestCase):
    """Fixture construction and CLI plumbing shared by the contract tests."""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory(prefix="docx-review-")
        self.addCleanup(self._tmp.cleanup)
        self.tmpdir = Path(self._tmp.name)

    # -- CLI ---------------------------------------------------------------

    def run_cli(self, *args: object) -> str:
        env = dict(os.environ)
        env["PYTHONIOENCODING"] = "utf-8"
        result = subprocess.run(
            [sys.executable, "-B", *[str(a) for a in args]],
            cwd=str(REPO_ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            env=env,
        )
        self.assertEqual(
            result.returncode,
            0,
            msg=(
                f"command failed: {args}\n"
                f"stdout:\n{result.stdout}\n"
                f"stderr:\n{result.stderr}"
            ),
        )
        return result.stdout

    def inspect_payload(self, docx: Path) -> dict:
        stdout = self.run_cli(INSPECT_CLI, "inspect", "--in", docx, "--json")
        return json.loads(stdout)

    def extract_payloads(self, docx: Path) -> tuple[dict, dict]:
        out_v2 = self.tmpdir / f"{docx.stem}-v2.json"
        out_v1 = self.tmpdir / f"{docx.stem}-v1.json"
        self.run_cli(
            EXTRACT_CLI,
            "--in",
            docx,
            "--out",
            out_v2,
            "--rag-output",
            out_v1,
        )
        v2 = json.loads(out_v2.read_text(encoding="utf-8"))
        v1 = json.loads(out_v1.read_text(encoding="utf-8"))
        return v2, v1

    # -- review access -----------------------------------------------------

    def review_of(self, payload: dict) -> dict:
        self.assertIn("review", payload)
        review = payload["review"]
        self.assertIsInstance(review, dict)
        return review

    def list_of(self, container: dict, key: str) -> list:
        self.assertIn(key, container)
        value = container[key]
        self.assertIsInstance(value, list)
        return value

    def anchor_of(self, comment: dict) -> dict:
        self.assertIn("anchor", comment)
        anchor = comment["anchor"]
        self.assertIsInstance(anchor, dict)
        return anchor

    def comment_by_id(self, comments: list, comment_id: str) -> dict:
        matches = [c for c in comments if c.get("id") == comment_id]
        self.assertEqual(len(matches), 1, msg=f"comments: {comments}")
        return matches[0]

    def assert_missing_metadata(self, entry: dict, expected: list) -> None:
        self.assertIn("missing_metadata", entry)
        self.assertEqual(sorted(entry["missing_metadata"]), expected)

    # -- fixtures ----------------------------------------------------------

    def build_docx(
        self,
        name: str,
        body_xml: str,
        *,
        comments_xml: str | None = None,
        comments_extended_xml: str | None = None,
        header_body_xml: str | None = None,
        extra_parts: dict[str, bytes] | None = None,
        extra_content_types: dict[str, str] | None = None,
        extra_relationships: dict[str, tuple[str, str]] | None = None,
    ) -> Path:
        path = self.tmpdir / name
        doc = Document()
        doc.add_paragraph("seed")
        if header_body_xml is not None:
            doc.sections[0].header.paragraphs[0].text = "Header"
        doc.save(str(path))

        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}

        document_xml = parts["word/document.xml"].decode("utf-8")
        match = re.search(r"<w:sectPr[\s\S]*?</w:sectPr>", document_xml)
        sect_pr = match.group(0) if match else ""
        parts["word/document.xml"] = _document_xml(body_xml, sect_pr).encode("utf-8")

        content_type_entries: list[str] = []
        relationship_entries: list[str] = []

        if comments_xml is not None:
            parts["word/comments.xml"] = _comments_xml(comments_xml).encode("utf-8")
            content_type_entries.append(
                '<Override PartName="/word/comments.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.'
                'wordprocessingml.comments+xml"/>'
            )
            relationship_entries.append(
                '<Relationship Id="rId900" Type="http://schemas.openxmlformats.org/'
                'officeDocument/2006/relationships/comments" Target="comments.xml"/>'
            )

        if comments_extended_xml is not None:
            parts["word/commentsExtended.xml"] = _comments_extended_xml(
                comments_extended_xml
            ).encode("utf-8")
            content_type_entries.append(
                '<Override PartName="/word/commentsExtended.xml" '
                'ContentType="application/vnd.ms-word.commentsExtended+xml"/>'
            )
            relationship_entries.append(
                '<Relationship Id="rId901" Type="http://schemas.microsoft.com/'
                'office/2011/relationships/commentsExtended" '
                'Target="commentsExtended.xml"/>'
            )

        if header_body_xml is not None:
            parts["word/header1.xml"] = _header_xml(header_body_xml).encode("utf-8")

        content_types = parts["[Content_Types].xml"].decode("utf-8")
        if content_type_entries:
            content_types = content_types.replace(
                "</Types>", "".join(content_type_entries) + "</Types>"
            )
            parts["[Content_Types].xml"] = content_types.encode("utf-8")

        relationships = parts["word/_rels/document.xml.rels"].decode("utf-8")
        if relationship_entries:
            relationships = relationships.replace(
                "</Relationships>", "".join(relationship_entries) + "</Relationships>"
            )
            parts["word/_rels/document.xml.rels"] = relationships.encode("utf-8")

        for part_name, data in (extra_parts or {}).items():
            parts[part_name] = data

        if extra_content_types:
            content_types = parts["[Content_Types].xml"].decode("utf-8")
            overrides = "".join(
                f'<Override PartName="/{part_name}" ContentType="{content_type}"/>'
                for part_name, content_type in extra_content_types.items()
            )
            parts["[Content_Types].xml"] = content_types.replace(
                "</Types>", overrides + "</Types>"
            ).encode("utf-8")

        if extra_relationships:
            relationships = parts["word/_rels/document.xml.rels"].decode("utf-8")
            entries = "".join(
                f'<Relationship Id="{relationship_id}" Type="{relationship_type}" '
                f'Target="{target}"/>'
                for relationship_id, (relationship_type, target) in (
                    extra_relationships.items()
                )
            )
            parts["word/_rels/document.xml.rels"] = relationships.replace(
                "</Relationships>", entries + "</Relationships>"
            ).encode("utf-8")

        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for part_name, data in parts.items():
                archive.writestr(part_name, data)
        return path


def _document_xml(body: str, sect_pr: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W_NS}" xmlns:w14="{W14_NS}" '
        f'xmlns:w15="{W15_NS}" xmlns:r="{R_NS}" xmlns:v="{V_NS}">'
        f"<w:body>{body}{sect_pr}</w:body></w:document>"
    )


def _comments_xml(body: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:comments xmlns:w="{W_NS}" xmlns:w14="{W14_NS}" '
        f'xmlns:w15="{W15_NS}">{body}</w:comments>'
    )


def _comments_extended_xml(body: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w15:commentsEx xmlns:w15="{W15_NS}">{body}</w15:commentsEx>'
    )


def _header_xml(body: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:hdr xmlns:w="{W_NS}" xmlns:w14="{W14_NS}" '
        f'xmlns:w15="{W15_NS}">{body}</w:hdr>'
    )


class ReviewSchemaTests(DocxReviewTestBase):
    """Top-level schema, text views and revision records."""

    def test_inspect_json_adds_review_and_keeps_legacy_view(self) -> None:
        docx = self.build_docx(
            "mixed.docx",
            MIXED_BODY,
            comments_xml=MIXED_COMMENTS,
            comments_extended_xml=MIXED_COMMENTS_EXTENDED,
        )
        payload = self.inspect_payload(docx)

        self.assertEqual(
            set(payload), {"comments", "revisions", "highlights", "review"}
        )
        self.assertEqual(
            payload["comments"],
            [
                {
                    "author": "Alice",
                    "date": "2026-10-01T10:00:00Z",
                    "text": "Parent note",
                }
            ],
        )
        self.assertEqual(
            payload["revisions"],
            {
                "insertions": [
                    {
                        "author": "Alice",
                        "date": "2026-10-01T10:00:00Z",
                        "text": "BC",
                    },
                    {"author": "Carol", "date": "Unknown", "text": "Nur"},
                ],
                "deletions": [{"text": "E"}, {"text": "F"}],
            },
        )
        self.assertEqual(payload["highlights"], [])

        review = self.review_of(payload)
        self.assertEqual(review["schema"], REVIEW_SCHEMA)
        self.assertEqual(
            set(review),
            {
                "schema",
                "paragraphs",
                "revisions",
                "comments",
                "limitations",
                "coverage",
            },
        )
        self.assertEqual(
            review["coverage"],
            {
                "supported_parts": [DOCUMENT_PART],
                "text_views_complete": True,
            },
        )
        self.assertEqual(self.list_of(review, "limitations"), [])

    def test_paragraph_text_views_and_locations_are_exact(self) -> None:
        docx = self.build_docx(
            "mixed.docx",
            MIXED_BODY,
            comments_xml=MIXED_COMMENTS,
            comments_extended_xml=MIXED_COMMENTS_EXTENDED,
        )
        review = self.review_of(self.inspect_payload(docx))
        paragraphs = self.list_of(review, "paragraphs")

        self.assertEqual(len(paragraphs), 4)
        self.assertEqual(
            paragraphs,
            [
                {
                    "location": _location(P1),
                    "text_original": "Intro",
                    "text_proposed": "Intro",
                    "text_markup": "Intro",
                },
                MIXED_PARAGRAPH_2,
                {
                    "location": _location(P3),
                    "text_original": "",
                    "text_proposed": "",
                    "text_markup": "",
                },
                {
                    "location": _location(P4),
                    "text_original": "",
                    "text_proposed": "Nur",
                    "text_markup": "Nur",
                },
            ],
        )
        for paragraph in paragraphs:
            self.assertEqual(
                set(paragraph),
                {
                    "location",
                    "text_original",
                    "text_proposed",
                    "text_markup",
                },
            )

    def test_insertion_revision_merges_runs_and_reports_ranges(self) -> None:
        docx = self.build_docx("insert.docx", MIXED_BODY)
        review = self.review_of(self.inspect_payload(docx))
        revisions = self.list_of(review, "revisions")

        self.assertEqual(len(revisions), 3)
        self.assertEqual(revisions[0], MIXED_REVISION_INSERT)
        self.assertEqual(
            revisions[0]["ranges"],
            {"original": [1, 1], "proposed": [1, 3], "markup": [1, 3]},
        )
        self.assertEqual(revisions[0]["text"], "BC")

    def test_deletion_revision_reports_collapsed_proposed_range(self) -> None:
        docx = self.build_docx("delete.docx", MIXED_BODY)
        review = self.review_of(self.inspect_payload(docx))
        revisions = self.list_of(review, "revisions")

        self.assertEqual(len(revisions), 3)
        self.assertEqual(revisions[1], MIXED_REVISION_DELETE)
        self.assertEqual(revisions[2], MIXED_REVISION_DATE_MISSING)
        self.assert_missing_metadata(revisions[2], ["date"])
        self.assertEqual(revisions[1]["text"], "EF")

    def test_deletion_ranges_use_true_text_boundary(self) -> None:
        """A deleted span collapses at the retained text limit, never one left."""

        cases = {
            "start": (
                '<w:del w:id="20" w:author="Alice" '
                'w:date="2026-10-01T10:00:00Z">'
                "<w:r><w:delText>AB</w:delText></w:r></w:del>"
                "<w:r><w:t>CD</w:t></w:r>",
                "ABCD",
                "CD",
                "ABCD",
                {"original": [0, 2], "proposed": [0, 0], "markup": [0, 2]},
            ),
            "middle": (
                "<w:r><w:t>A</w:t></w:r>"
                '<w:del w:id="21" w:author="Alice" '
                'w:date="2026-10-01T10:00:00Z">'
                "<w:r><w:delText>B</w:delText></w:r></w:del>"
                "<w:r><w:t>C</w:t></w:r>",
                "ABC",
                "AC",
                "ABC",
                {"original": [1, 2], "proposed": [1, 1], "markup": [1, 2]},
            ),
            "end": (
                "<w:r><w:t>ABC</w:t></w:r>"
                '<w:del w:id="22" w:author="Alice" '
                'w:date="2026-10-01T10:00:00Z">'
                "<w:r><w:delText>D</w:delText></w:r></w:del>",
                "ABCD",
                "ABC",
                "ABCD",
                {"original": [3, 4], "proposed": [3, 3], "markup": [3, 4]},
            ),
            "proposed_text_before_deletion": (
                "<w:r><w:t>ABCD</w:t></w:r>"
                '<w:del w:id="23" w:author="Bob" '
                'w:date="2026-10-02T11:30:00Z">'
                "<w:r><w:delText>EF</w:delText></w:r></w:del>",
                "ABCDEF",
                "ABCD",
                "ABCDEF",
                {"original": [4, 6], "proposed": [4, 4], "markup": [4, 6]},
            ),
        }

        for case_id, (
            body,
            text_original,
            text_proposed,
            text_markup,
            ranges,
        ) in cases.items():
            with self.subTest(case=case_id):
                docx = self.build_docx(
                    f"boundary-{case_id}.docx", f"<w:p>{body}</w:p>"
                )
                review = self.review_of(self.inspect_payload(docx))
                paragraphs = self.list_of(review, "paragraphs")

                self.assertEqual(len(paragraphs), 1)
                self.assertEqual(
                    paragraphs[0],
                    {
                        "location": _location(P1),
                        "text_original": text_original,
                        "text_proposed": text_proposed,
                        "text_markup": text_markup,
                    },
                )

                revisions = self.list_of(review, "revisions")
                self.assertEqual(len(revisions), 1)
                self.assertEqual(revisions[0]["kind"], "delete")
                self.assertEqual(revisions[0]["ranges"], ranges)

    def test_whitespace_tab_break_and_multibyte_are_preserved(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:r><w:t>\u00e4</w:t><w:tab/><w:t>B</w:t><w:br/><w:t>C</w:t></w:r>
            <w:ins w:id="60" w:author="Alice" w:date="2026-10-01T10:00:00Z"><w:r><w:t>XYZ</w:t></w:r></w:ins></w:p>
            <w:p><w:r><w:t xml:space="preserve">  A  </w:t></w:r></w:p>
            """
        )
        docx = self.build_docx("whitespace.docx", body)
        review = self.review_of(self.inspect_payload(docx))
        paragraphs = self.list_of(review, "paragraphs")

        self.assertEqual(len(paragraphs), 2)
        self.assertEqual(paragraphs[0]["text_markup"], "\u00e4\tB\nCXYZ")
        self.assertEqual(paragraphs[0]["text_original"], "\u00e4\tB\nC")
        self.assertEqual(paragraphs[0]["text_proposed"], "\u00e4\tB\nCXYZ")
        self.assertEqual(len(paragraphs[0]["text_markup"]), 8)
        self.assertEqual(paragraphs[1]["text_markup"], "  A  ")
        self.assertEqual(paragraphs[1]["text_original"], "  A  ")
        self.assertEqual(paragraphs[1]["text_proposed"], "  A  ")

        revisions = self.list_of(review, "revisions")
        self.assertEqual(len(revisions), 1)
        self.assertEqual(revisions[0]["text"], "XYZ")
        self.assertEqual(
            revisions[0]["ranges"],
            {"original": [5, 5], "proposed": [5, 8], "markup": [5, 8]},
        )

    def test_plain_document_has_empty_review_and_complete_coverage(self) -> None:
        docx = self.build_docx(
            "plain.docx",
            "<w:p><w:r><w:t>Plain text</w:t></w:r></w:p><w:p/>",
        )
        review = self.review_of(self.inspect_payload(docx))

        self.assertEqual(review["schema"], REVIEW_SCHEMA)
        self.assertEqual(
            self.list_of(review, "paragraphs"),
            [
                {
                    "location": _location(P1),
                    "text_original": "Plain text",
                    "text_proposed": "Plain text",
                    "text_markup": "Plain text",
                },
                {
                    "location": _location(P2),
                    "text_original": "",
                    "text_proposed": "",
                    "text_markup": "",
                },
            ],
        )
        self.assertEqual(self.list_of(review, "revisions"), [])
        self.assertEqual(self.list_of(review, "comments"), [])
        self.assertEqual(self.list_of(review, "limitations"), [])
        self.assertEqual(
            review["coverage"],
            {
                "supported_parts": [DOCUMENT_PART],
                "text_views_complete": True,
            },
        )


class CommentAnchorTests(DocxReviewTestBase):
    """Comment anchoring, segments, threads and diagnostics."""

    def test_comment_anchor_targets_inserted_text(self) -> None:
        docx = self.build_docx(
            "anchor-insert.docx",
            MIXED_BODY,
            comments_xml=MIXED_COMMENTS,
            comments_extended_xml=MIXED_COMMENTS_EXTENDED,
        )
        review = self.review_of(self.inspect_payload(docx))
        comments = self.list_of(review, "comments")

        self.assertEqual(len(comments), 1)
        comment = self.comment_by_id(comments, "1")
        self.assertEqual(
            {key: value for key, value in comment.items() if key != "anchor"},
            {
                "id": "1",
                "author": "Alice",
                "date": "2026-10-01T10:00:00Z",
                "text": "Parent note",
                "missing_metadata": [],
                "resolved": True,
                "parent_id": None,
                "reply_ids": [],
            },
        )

        anchor = self.anchor_of(comment)
        self.assertEqual(anchor["status"], "resolved")
        self.assertEqual(anchor["start"], {"location": _location(P2), "offset": 1})
        self.assertEqual(anchor["end"], {"location": _location(P2), "offset": 3})
        self.assertEqual(
            anchor["segments"],
            [
                {
                    "location": _location(P2),
                    "ranges": {
                        "original": [1, 1],
                        "proposed": [1, 3],
                        "markup": [1, 3],
                    },
                    "text_original": "",
                    "text_proposed": "BC",
                    "text_markup": "BC",
                }
            ],
        )
        self.assertEqual(anchor.get("diagnostics", []), [])

    def test_comment_text_joins_paragraphs_and_drops_trailing_empty(self) -> None:
        body = (
            "<w:p><w:commentRangeStart w:id=\"1\"/>"
            "<w:r><w:t>Intro</w:t></w:r>"
            "<w:commentRangeEnd w:id=\"1\"/>"
            "<w:r><w:commentReference w:id=\"1\"/></w:r></w:p>"
        )
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1" w:author="Alice" w:date="2026-10-01T10:00:00Z">
              <w:p w14:paraId="AAAA1111"><w:r><w:t>First line</w:t></w:r></w:p>
              <w:p><w:r><w:t>Second line</w:t></w:r></w:p>
              <w:p/>
            </w:comment>
            """
        )
        docx = self.build_docx("comment-text.docx", body, comments_xml=comments)
        review = self.review_of(self.inspect_payload(docx))
        comment = self.comment_by_id(self.list_of(review, "comments"), "1")

        self.assertEqual(comment["text"], "First line\nSecond line")
        self.assertEqual(
            self.anchor_of(comment)["segments"][0]["text_markup"], "Intro"
        )

    def test_cross_paragraph_comment_keeps_all_segments(self) -> None:
        body = (
            "<w:p><w:r><w:t>In</w:t></w:r>"
            "<w:commentRangeStart w:id=\"1\"/>"
            "<w:r><w:t>tro</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>ABC</w:t></w:r>"
            "<w:commentRangeEnd w:id=\"1\"/>"
            "<w:r><w:commentReference w:id=\"1\"/></w:r></w:p>"
        )
        comments = (
            '<w:comment w:id="1" w:author="Alice" '
            'w:date="2026-10-01T10:00:00Z">'
            '<w:p w14:paraId="AAAA1111">'
            '<w:r><w:t>Across paragraphs</w:t></w:r></w:p>'
            "</w:comment>"
        )
        docx = self.build_docx("cross-paragraph.docx", body, comments_xml=comments)
        review = self.review_of(self.inspect_payload(docx))
        anchor = self.anchor_of(
            self.comment_by_id(self.list_of(review, "comments"), "1")
        )

        self.assertEqual(anchor["status"], "resolved")
        self.assertEqual(anchor["start"], {"location": _location(P1), "offset": 2})
        self.assertEqual(anchor["end"], {"location": _location(P2), "offset": 3})
        self.assertEqual(
            anchor["segments"],
            [
                {
                    "location": _location(P1),
                    "ranges": {
                        "original": [2, 5],
                        "proposed": [2, 5],
                        "markup": [2, 5],
                    },
                    "text_original": "tro",
                    "text_proposed": "tro",
                    "text_markup": "tro",
                },
                {
                    "location": _location(P2),
                    "ranges": {
                        "original": [0, 3],
                        "proposed": [0, 3],
                        "markup": [0, 3],
                    },
                    "text_original": "ABC",
                    "text_proposed": "ABC",
                    "text_markup": "ABC",
                },
            ],
        )

    def test_table_cell_comment_location(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:r><w:t>Before</w:t></w:r></w:p>
            <w:tbl><w:tr><w:tc><w:p>
              <w:commentRangeStart w:id="1"/><w:r><w:t>Cell</w:t></w:r>
              <w:commentRangeEnd w:id="1"/>
              <w:r><w:commentReference w:id="1"/></w:r>
            </w:p></w:tc></w:tr></w:tbl>
            """
        )
        comments = (
            '<w:comment w:id="1" w:author="Alice" '
            'w:date="2026-10-01T10:00:00Z">'
            '<w:p w14:paraId="AAAA1111">'
            '<w:r><w:t>Table note</w:t></w:r></w:p></w:comment>'
        )
        docx = self.build_docx("table-comment.docx", body, comments_xml=comments)
        review = self.review_of(self.inspect_payload(docx))

        paths = [p["location"]["path"] for p in review["paragraphs"]]
        self.assertEqual(paths, [P1, TABLE_CELL_P])

        anchor = self.anchor_of(
            self.comment_by_id(self.list_of(review, "comments"), "1")
        )
        self.assertEqual(anchor["status"], "resolved")
        self.assertEqual(
            anchor["start"], {"location": _location(TABLE_CELL_P), "offset": 0}
        )
        self.assertEqual(
            anchor["end"], {"location": _location(TABLE_CELL_P), "offset": 4}
        )
        self.assertEqual(anchor["segments"][0]["text_markup"], "Cell")

    def test_nested_table_comment_location(self) -> None:
        body = textwrap.dedent(
            """\
            <w:tbl><w:tr><w:tc>
              <w:p><w:r><w:t>Outer</w:t></w:r></w:p>
              <w:tbl><w:tr><w:tc><w:p>
                <w:commentRangeStart w:id="1"/><w:r><w:t>Inner</w:t></w:r>
                <w:commentRangeEnd w:id="1"/>
                <w:r><w:commentReference w:id="1"/></w:r>
              </w:p></w:tc></w:tr></w:tbl>
            </w:tc></w:tr></w:tbl>
            """
        )
        comments = (
            '<w:comment w:id="1" w:author="Alice" '
            'w:date="2026-10-01T10:00:00Z">'
            '<w:p w14:paraId="AAAA1111">'
            '<w:r><w:t>Nested note</w:t></w:r></w:p></w:comment>'
        )
        docx = self.build_docx("nested-table.docx", body, comments_xml=comments)
        review = self.review_of(self.inspect_payload(docx))

        paths = [p["location"]["path"] for p in review["paragraphs"]]
        self.assertIn(NESTED_TABLE_CELL_P, paths)

        anchor = self.anchor_of(
            self.comment_by_id(self.list_of(review, "comments"), "1")
        )
        self.assertEqual(
            anchor["start"],
            {"location": _location(NESTED_TABLE_CELL_P), "offset": 0},
        )
        self.assertEqual(
            anchor["end"],
            {"location": _location(NESTED_TABLE_CELL_P), "offset": 5},
        )
        self.assertEqual(anchor["segments"][0]["text_markup"], "Inner")

    def test_comment_status_reply_thread_and_missing_metadata(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:commentRangeStart w:id="1"/><w:r><w:t>I</w:t></w:r>
            <w:commentRangeStart w:id="3"/><w:r><w:t>n</w:t></w:r>
            <w:commentRangeEnd w:id="1"/><w:r><w:t>tro</w:t></w:r>
            <w:commentRangeEnd w:id="3"/>
            <w:r><w:commentReference w:id="1"/><w:commentReference w:id="3"/></w:r></w:p>
            <w:p><w:commentRangeStart w:id="4"/><w:r><w:t>More</w:t></w:r>
            <w:commentRangeEnd w:id="4"/>
            <w:r><w:commentReference w:id="4"/></w:r></w:p>
            """
        )
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1" w:author="Alice" w:date="2026-10-01T10:00:00Z">
              <w:p w14:paraId="AAAA1111"><w:r><w:t>Parent comment</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="2" w:author="Bob" w:date="2026-10-02T11:30:00Z">
              <w:p w14:paraId="BBBB2222"><w:r><w:t>Reply text</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="3" w:author="Carol" w:date="2026-10-03T12:00:00Z">
              <w:p w14:paraId="CCCC3333"><w:r><w:t>Done false</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="4">
              <w:p w14:paraId="DDDD4444"><w:r><w:t>No metadata</w:t></w:r></w:p>
            </w:comment>
            """
        )
        comments_extended = textwrap.dedent(
            """\
            <w15:commentEx w15:paraId="AAAA1111" w15:done="1"/>
            <w15:commentEx w15:paraId="BBBB2222" w15:paraIdParent="AAAA1111"/>
            <w15:commentEx w15:paraId="CCCC3333" w15:done="0"/>
            """
        )
        docx = self.build_docx(
            "threads.docx",
            body,
            comments_xml=comments,
            comments_extended_xml=comments_extended,
        )
        review = self.review_of(self.inspect_payload(docx))
        comments_out = self.list_of(review, "comments")
        self.assertEqual(len(comments_out), 4)

        parent = self.comment_by_id(comments_out, "1")
        self.assertEqual(parent["text"], "Parent comment")
        self.assertEqual(parent["resolved"], True)
        self.assertEqual(parent["parent_id"], None)
        self.assertEqual(parent["reply_ids"], ["2"])

        reply = self.comment_by_id(comments_out, "2")
        self.assertEqual(reply["text"], "Reply text")
        self.assertEqual(reply["resolved"], None)
        self.assertEqual(reply["parent_id"], "1")
        self.assertEqual(reply["reply_ids"], [])
        reply_anchor = self.anchor_of(reply)
        self.assertEqual(reply_anchor["status"], "thread_only")
        self.assertIsNone(reply_anchor["start"])
        self.assertIsNone(reply_anchor["end"])
        self.assertEqual(reply_anchor["segments"], [])

        done_false = self.comment_by_id(comments_out, "3")
        self.assertEqual(done_false["resolved"], False)
        self.assertEqual(done_false["parent_id"], None)
        self.assertEqual(done_false["reply_ids"], [])

        # Comment 1 covers markup [0, 2), comment 3 covers [1, 5): they overlap.
        parent_anchor = self.anchor_of(parent)
        self.assertEqual(parent_anchor["start"], {"location": _location(P1), "offset": 0})
        self.assertEqual(parent_anchor["end"], {"location": _location(P1), "offset": 2})
        overlap_anchor = self.anchor_of(done_false)
        self.assertEqual(overlap_anchor["start"], {"location": _location(P1), "offset": 1})
        self.assertEqual(overlap_anchor["end"], {"location": _location(P1), "offset": 5})

        missing = self.comment_by_id(comments_out, "4")
        self.assertIsNone(missing["author"])
        self.assertIsNone(missing["date"])
        self.assertIsNone(missing["resolved"])
        self.assert_missing_metadata(missing, ["author", "date"])

    def test_unresolved_anchor_diagnostics(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:commentRangeEnd w:id="1"/>
            <w:r><w:t>One</w:t></w:r>
            <w:r><w:commentReference w:id="1"/></w:r></w:p>
            <w:p><w:commentRangeStart w:id="2"/><w:r><w:t>Dup</w:t></w:r>
            <w:commentRangeStart w:id="2"/><w:r><w:t>More</w:t></w:r>
            <w:commentRangeEnd w:id="2"/>
            <w:r><w:commentReference w:id="2"/></w:r></w:p>
            <w:p><w:commentRangeEnd w:id="3"/><w:r><w:t>Rev</w:t></w:r>
            <w:commentRangeStart w:id="3"/>
            <w:r><w:commentReference w:id="3"/></w:r></w:p>
            """
        )
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1" w:author="Alice" w:date="2026-10-01T10:00:00Z">
              <w:p w14:paraId="AAAA1111"><w:r><w:t>Orphan</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="2" w:author="Bob" w:date="2026-10-02T11:30:00Z">
              <w:p w14:paraId="BBBB2222"><w:r><w:t>Duplicate</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="3" w:author="Carol" w:date="2026-10-03T12:00:00Z">
              <w:p w14:paraId="CCCC3333"><w:r><w:t>Reversed</w:t></w:r></w:p>
            </w:comment>
            """
        )
        docx = self.build_docx("broken-anchors.docx", body, comments_xml=comments)
        review = self.review_of(self.inspect_payload(docx))
        comments_out = self.list_of(review, "comments")

        cases = {
            "1": "orphan_anchor",
            "2": "duplicate_anchor",
            "3": "reversed_anchor",
        }
        for comment_id, expected_code in cases.items():
            with self.subTest(comment_id=comment_id):
                anchor = self.anchor_of(
                    self.comment_by_id(comments_out, comment_id)
                )
                self.assertEqual(anchor["status"], "unresolved")
                self.assertIn("diagnostics", anchor)
                codes = [item.get("code") for item in anchor["diagnostics"]]
                self.assertEqual(codes, [expected_code])
                self.assertIsNone(anchor["start"])
                self.assertIsNone(anchor["end"])
                self.assertEqual(anchor["segments"], [])

    def test_duplicate_comment_ids_never_resolve_anchors(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:commentRangeStart w:id="1"/><w:r><w:t>Alpha</w:t></w:r>
            <w:commentRangeEnd w:id="1"/>
            <w:r><w:commentReference w:id="1"/></w:r></w:p>
            <w:p><w:commentRangeStart w:id="5"/><w:r><w:t>Beta</w:t></w:r>
            <w:commentRangeEnd w:id="5"/>
            <w:r><w:commentReference w:id="5"/></w:r></w:p>
            """
        )
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1" w:author="Alice" w:date="2026-10-01T10:00:00Z">
              <w:p w14:paraId="AAAA1111"><w:r><w:t>Unique</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="5" w:author="Bob" w:date="2026-10-02T11:30:00Z">
              <w:p w14:paraId="BBBB2222"><w:r><w:t>First copy</w:t></w:r></w:p>
            </w:comment>
            <w:comment w:id="5" w:author="Carol" w:date="2026-10-03T12:00:00Z">
              <w:p w14:paraId="CCCC3333"><w:r><w:t>Second copy</w:t></w:r></w:p>
            </w:comment>
            """
        )
        docx = self.build_docx("duplicate-ids.docx", body, comments_xml=comments)
        review = self.review_of(self.inspect_payload(docx))
        comments_out = self.list_of(review, "comments")

        self.assertEqual(len(comments_out), 3)
        for index in (1, 2):
            with self.subTest(comment_index=index):
                duplicate = comments_out[index]
                self.assertEqual(duplicate["id"], "5")
                anchor = self.anchor_of(duplicate)
                self.assertEqual(anchor["status"], "unresolved")
                self.assertIsNone(anchor["start"])
                self.assertIsNone(anchor["end"])
                self.assertEqual(anchor["segments"], [])
                self.assertEqual(
                    [item.get("code") for item in anchor["diagnostics"]],
                    ["ambiguous_comment_id"],
                )

        ambiguous = [
            item
            for item in self.list_of(review, "limitations")
            if item["code"] == "ambiguous_comment_id"
        ]
        self.assertEqual(
            {item["location"]["path"] for item in ambiguous},
            {
                "/comments[1]/comment[2]",
                "/comments[1]/comment[3]",
            },
        )
        self.assertEqual(
            {item["details"]["comment_id"] for item in ambiguous}, {"5"}
        )

    def test_comment_ex_done_accepts_boolean_literals(self) -> None:
        body_parts = []
        for comment_id in ("1", "2", "3", "4"):
            body_parts.append(
                f'<w:p><w:commentRangeStart w:id="{comment_id}"/>'
                f"<w:r><w:t>Text {comment_id}</w:t></w:r>"
                f'<w:commentRangeEnd w:id="{comment_id}"/>'
                f'<w:r><w:commentReference w:id="{comment_id}"/></w:r></w:p>'
            )
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1"><w:p w14:paraId="AAAA1111"><w:r><w:t>One</w:t></w:r></w:p></w:comment>
            <w:comment w:id="2"><w:p w14:paraId="BBBB2222"><w:r><w:t>Two</w:t></w:r></w:p></w:comment>
            <w:comment w:id="3"><w:p w14:paraId="CCCC3333"><w:r><w:t>Three</w:t></w:r></w:p></w:comment>
            <w:comment w:id="4"><w:p w14:paraId="DDDD4444"><w:r><w:t>Four</w:t></w:r></w:p></w:comment>
            """
        )
        comments_extended = textwrap.dedent(
            """\
            <w15:commentEx w15:paraId="AAAA1111" w15:done="true"/>
            <w15:commentEx w15:paraId="BBBB2222" w15:done="bogus"/>
            <w15:commentEx w15:paraId="CCCC3333" w15:done="false"/>
            <w15:commentEx w15:paraId="DDDD4444"/>
            """
        )
        docx = self.build_docx(
            "done-boolean.docx",
            "".join(body_parts),
            comments_xml=comments,
            comments_extended_xml=comments_extended,
        )
        review = self.review_of(self.inspect_payload(docx))
        comments_out = self.list_of(review, "comments")

        self.assertEqual(
            [self.comment_by_id(comments_out, item)["resolved"] for item in "1234"],
            [True, None, False, None],
        )

        malformed = [
            item
            for item in self.list_of(review, "limitations")
            if item["code"] == "malformed_thread_metadata"
        ]
        self.assertIn(
            "/commentsEx[1]/commentEx[2]",
            {item["location"]["path"] for item in malformed},
        )
        self.assertEqual(
            {
                item["details"].get("comment_id")
                for item in malformed
                if item["location"]["path"] == "/commentsEx[1]/commentEx[2]"
            },
            {"2"},
        )
        self.assertNotIn(
            "1",
            {item["details"].get("comment_id") for item in malformed},
        )
        self.assertNotIn(
            "3",
            {item["details"].get("comment_id") for item in malformed},
        )

    def test_duplicate_last_para_id_does_not_bind_thread_metadata(self) -> None:
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1"><w:p w14:paraId="DUP00001"><w:r><w:t>One</w:t></w:r></w:p></w:comment>
            <w:comment w:id="2"><w:p w14:paraId="DUP00001"><w:r><w:t>Two</w:t></w:r></w:p></w:comment>
            <w:comment w:id="3"><w:p w14:paraId="PARENT01"><w:r><w:t>Three</w:t></w:r></w:p></w:comment>
            """
        )
        comments_extended = textwrap.dedent(
            """\
            <w15:commentEx w15:paraId="DUP00001" w15:done="true" w15:paraIdParent="PARENT01"/>
            <w15:commentEx w15:paraId="PARENT01"/>
            """
        )
        body = (
            '<w:p><w:commentRangeStart w:id="1"/><w:r><w:t>One</w:t></w:r>'
            '<w:commentRangeEnd w:id="1"/><w:r><w:commentReference w:id="1"/></w:r></w:p>'
            '<w:p><w:commentRangeStart w:id="2"/><w:r><w:t>Two</w:t></w:r>'
            '<w:commentRangeEnd w:id="2"/><w:r><w:commentReference w:id="2"/></w:r></w:p>'
        )
        docx = self.build_docx(
            "duplicate-para-id.docx",
            body,
            comments_xml=comments,
            comments_extended_xml=comments_extended,
        )
        review = self.review_of(self.inspect_payload(docx))
        comments_out = self.list_of(review, "comments")

        for comment_id in ("1", "2"):
            with self.subTest(comment_id=comment_id):
                comment = self.comment_by_id(comments_out, comment_id)
                self.assertIsNone(comment["resolved"])
                self.assertIsNone(comment["parent_id"])
                self.assertEqual(comment["reply_ids"], [])
        self.assertEqual(
            self.comment_by_id(comments_out, "3")["reply_ids"], []
        )

        diagnostics = [
            item["code"] for item in self.list_of(review, "limitations")
        ]
        self.assertTrue(
            any(
                "ambiguous" in code or "thread" in code or "parent" in code
                for code in diagnostics
            ),
            msg=diagnostics,
        )

    def test_self_and_cyclic_parent_relations_are_diagnosed(self) -> None:
        comments = textwrap.dedent(
            """\
            <w:comment w:id="1"><w:p w14:paraId="SELF0001"><w:r><w:t>Self</w:t></w:r></w:p></w:comment>
            <w:comment w:id="2"><w:p w14:paraId="CYCLE002"><w:r><w:t>Cycle A</w:t></w:r></w:p></w:comment>
            <w:comment w:id="3"><w:p w14:paraId="CYCLE003"><w:r><w:t>Cycle B</w:t></w:r></w:p></w:comment>
            """
        )
        comments_extended = textwrap.dedent(
            """\
            <w15:commentEx w15:paraId="SELF0001" w15:paraIdParent="SELF0001"/>
            <w15:commentEx w15:paraId="CYCLE002" w15:paraIdParent="CYCLE003"/>
            <w15:commentEx w15:paraId="CYCLE003" w15:paraIdParent="CYCLE002"/>
            """
        )
        body = "".join(
            f'<w:p><w:commentRangeStart w:id="{comment_id}"/>'
            f"<w:r><w:t>Text {comment_id}</w:t></w:r>"
            f'<w:commentRangeEnd w:id="{comment_id}"/>'
            f'<w:r><w:commentReference w:id="{comment_id}"/></w:r></w:p>'
            for comment_id in ("1", "2", "3")
        )
        docx = self.build_docx(
            "cyclic-parent.docx",
            body,
            comments_xml=comments,
            comments_extended_xml=comments_extended,
        )
        review = self.review_of(self.inspect_payload(docx))
        comments_out = self.list_of(review, "comments")

        for comment in comments_out:
            with self.subTest(comment_id=comment["id"]):
                self.assertIsNone(comment["parent_id"])
                self.assertEqual(comment["reply_ids"], [])

        diagnostics = [
            item["code"] for item in self.list_of(review, "limitations")
        ]
        self.assertTrue(
            any(
                "parent" in code or "cycle" in code for code in diagnostics
            ),
            msg=diagnostics,
        )


class UnsupportedRevisionTests(DocxReviewTestBase):
    """Detected but unsupported revisions must surface as limitations."""

    def test_comment_markers_without_id_are_diagnosed_at_exact_paths(self) -> None:
        body = (
            '<w:p><w:commentRangeStart/><w:r><w:t>Text</w:t></w:r>'
            '<w:commentRangeEnd/><w:r><w:commentReference/></w:r></w:p>'
        )
        docx = self.build_docx("missing-marker-ids.docx", body)
        review = self.review_of(self.inspect_payload(docx))
        diagnostics = [
            item for item in review["limitations"]
            if item["code"] == "missing_marker_id"
        ]
        self.assertEqual(len(diagnostics), 3)
        self.assertEqual(
            {item["location"]["path"] for item in diagnostics},
            {
                P1 + "/commentRangeStart[1]",
                P1 + "/commentRangeEnd[1]",
                P1 + "/r[2]/commentReference[1]",
            },
        )
        self.assertTrue(all(item["part"] == DOCUMENT_PART for item in diagnostics))
        self.assertEqual(review["comments"], [])
        self.assertEqual(review["paragraphs"][0]["text_markup"], "Text")

    def test_structural_move_and_property_revisions_are_limitations(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:pPr><w:pPrChange w:id="30" w:author="Alice" w:date="2026-10-01T10:00:00Z"><w:pPr/></w:pPrChange></w:pPr>
              <w:r><w:rPr><w:rPrChange w:id="31" w:author="Alice" w:date="2026-10-01T10:00:00Z"><w:rPr/></w:rPrChange></w:rPr><w:t>Changed</w:t></w:r>
              <w:moveFrom w:id="32" w:author="Bob" w:date="2026-10-02T11:30:00Z"><w:r><w:t>Moved</w:t></w:r></w:moveFrom>
              <w:moveTo w:id="32" w:author="Bob" w:date="2026-10-02T11:30:00Z"><w:r><w:t>Moved</w:t></w:r></w:moveTo>
            </w:p>
            <w:p><w:pPr><w:rPr><w:del w:id="34" w:author="Alice" w:date="2026-10-01T10:00:00Z"/></w:rPr></w:pPr><w:r><w:t>Merged</w:t></w:r></w:p>
            <w:tbl><w:tr><w:trPr><w:ins w:id="33" w:author="Alice" w:date="2026-10-01T10:00:00Z"/></w:trPr><w:tc><w:p><w:r><w:t>Row</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
            """
        )
        docx = self.build_docx("unsupported-main.docx", body)
        review = self.review_of(self.inspect_payload(docx))
        limitations = self.list_of(review, "limitations")

        self.assertEqual(
            {item["code"] for item in limitations},
            {
                "unsupported_structural_revision",
                "unsupported_move",
                "unsupported_property_change",
            },
        )
        for item in limitations:
            self.assertEqual(
                set(item), {"code", "part", "location", "details"}
            )
            self.assertEqual(item["part"], DOCUMENT_PART)
        self.assertEqual(
            review["coverage"],
            {
                "supported_parts": [DOCUMENT_PART],
                "text_views_complete": False,
            },
        )

    def test_revisions_outside_main_story_are_limitations(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:r><w:t>Body</w:t></w:r>
              <w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p>
                <w:ins w:id="41" w:author="Alice" w:date="2026-10-01T10:00:00Z"><w:r><w:t>Box change</w:t></w:r></w:ins>
              </w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>
            """
        )
        header = (
            '<w:p><w:ins w:id="40" w:author="Alice" '
            'w:date="2026-10-01T10:00:00Z">'
            "<w:r><w:t>Header change</w:t></w:r></w:ins></w:p>"
        )
        docx = self.build_docx(
            "unsupported-story.docx",
            body,
            header_body_xml=header,
        )
        review = self.review_of(self.inspect_payload(docx))
        limitations = self.list_of(review, "limitations")
        story = [item for item in limitations if item["code"] == "unsupported_story"]

        self.assertEqual(
            {item["part"] for item in story},
            {"word/header1.xml", DOCUMENT_PART},
        )
        self.assertTrue(
            any(
                "txbxContent" in json.dumps(item["location"])
                for item in story
                if item["part"] == DOCUMENT_PART
            ),
            msg=f"limitations: {limitations}",
        )
        self.assertEqual(review["coverage"]["text_views_complete"], False)

    def test_review_markers_in_comments_and_glossary_are_stories(self) -> None:
        comments = (
            '<w:comment w:id="1" w:author="Alice" '
            'w:date="2026-10-01T10:00:00Z"><w:p>'
            '<w:ins w:id="70" w:author="Alice" '
            'w:date="2026-10-01T10:00:00Z">'
            "<w:r><w:t>Inserted note</w:t></w:r></w:ins></w:p></w:comment>"
        )
        glossary = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:glossaryDocument xmlns:w="{W_NS}"><w:docParts><w:docPart>'
            "<w:docPartBody><w:p>"
            '<w:del w:id="71" w:author="Bob" '
            'w:date="2026-10-02T11:30:00Z">'
            "<w:r><w:delText>Glossary old</w:delText></w:r></w:del>"
            "</w:p></w:docPartBody></w:docPart></w:docParts>"
            "</w:glossaryDocument>"
        )
        docx = self.build_docx(
            "other-stories.docx",
            "<w:p><w:r><w:t>Body</w:t></w:r></w:p>",
            comments_xml=comments,
            extra_parts={"word/glossary/document.xml": glossary.encode("utf-8")},
            extra_content_types={
                "word/glossary/document.xml": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document.glossary+xml"
                )
            },
            extra_relationships={
                "rId902": (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/glossaryDocument",
                    "glossary/document.xml",
                )
            },
        )
        review = self.review_of(self.inspect_payload(docx))

        story = [
            item
            for item in self.list_of(review, "limitations")
            if item["code"] == "unsupported_story"
        ]
        self.assertEqual(
            {item["part"] for item in story},
            {"word/comments.xml", "word/glossary/document.xml"},
        )
        self.assertEqual(review["coverage"]["text_views_complete"], False)

    def test_body_custom_xml_is_extracted_or_diagnosed(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:r><w:t>Before</w:t></w:r></w:p>
            <w:customXml w:element="wrapper">
              <w:p><w:ins w:id="80" w:author="Alice" w:date="2026-10-01T10:00:00Z">
                <w:r><w:t>Inside wrapper</w:t></w:r></w:ins></w:p>
            </w:customXml>
            """
        )
        docx = self.build_docx("custom-xml.docx", body)
        review = self.review_of(self.inspect_payload(docx))

        extracted = any(
            "Inside wrapper" in paragraph["text_markup"]
            for paragraph in self.list_of(review, "paragraphs")
        )
        diagnosed = any(
            "customXml" in json.dumps(item)
            for item in self.list_of(review, "limitations")
        )

        self.assertTrue(
            extracted or diagnosed,
            msg=f"paragraphs={review['paragraphs']} limitations={review['limitations']}",
        )
        if not extracted:
            self.assertFalse(review["coverage"]["text_views_complete"])

    def test_foreign_namespace_lookalikes_are_not_word_revisions(self) -> None:
        body = textwrap.dedent(
            """\
            <w:p><w:r><w:t>Keep</w:t></w:r>
              <x:ins xmlns:x="urn:example:foreign"><x:r><x:t>FakeIns</x:t></x:r></x:ins>
              <x:del xmlns:x="urn:example:foreign"><x:r><x:t>FakeDel</x:t></x:r></x:del>
              <x:t xmlns:x="urn:example:foreign">FakeText</x:t></w:p>
            <x:p xmlns:x="urn:example:foreign"><x:r><x:t>ForeignParagraph</x:t></x:r></x:p>
            """
        )
        docx = self.build_docx("foreign-namespace.docx", body)
        review = self.review_of(self.inspect_payload(docx))

        for paragraph in self.list_of(review, "paragraphs"):
            for view in ("text_original", "text_proposed", "text_markup"):
                self.assertNotIn("Fake", paragraph[view])
                self.assertNotIn("ForeignParagraph", paragraph[view])
        self.assertEqual(
            [
                revision["text"]
                for revision in self.list_of(review, "revisions")
            ],
            [],
        )
        self.assertTrue(
            self.list_of(review, "limitations"),
            msg="unknown foreign content must be reported, not silently ignored",
        )
        self.assertEqual(review["coverage"]["text_views_complete"], False)


class CompatibilityTests(DocxReviewTestBase):
    """The review is identical across surfaces and legacy payloads survive."""

    def test_inspect_v2_and_v1_share_identical_review(self) -> None:
        docx = self.build_docx(
            "mixed.docx",
            MIXED_BODY,
            comments_xml=MIXED_COMMENTS,
            comments_extended_xml=MIXED_COMMENTS_EXTENDED,
        )
        inspect_review = self.review_of(self.inspect_payload(docx))
        v2, v1 = self.extract_payloads(docx)
        v2_review = self.review_of(v2)
        v1_review = self.review_of(v1)

        self.assertEqual(v2_review, inspect_review)
        self.assertEqual(v1_review, inspect_review)
        self.assertEqual(v2_review["schema"], REVIEW_SCHEMA)
        self.assertEqual(len(self.list_of(v2_review, "revisions")), 3)
        self.assertEqual(len(self.list_of(v2_review, "comments")), 1)
        for block in v1["blocks"]:
            self.assertNotIn("review", block)
        self.assertNotIn("review", v2["document"])

    def test_legacy_defaults_remain_and_review_metadata_is_null(self) -> None:
        body = (
            "<w:p><w:commentRangeStart w:id=\"1\"/>"
            "<w:r><w:t>Text</w:t></w:r>"
            "<w:commentRangeEnd w:id=\"1\"/>"
            "<w:r><w:commentReference w:id=\"1\"/></w:r>"
            "<w:ins w:id=\"50\"><w:r><w:t>New</w:t></w:r></w:ins></w:p>"
        )
        comments = (
            '<w:comment w:id="1">'
            '<w:p w14:paraId="AAAA1111">'
            '<w:r><w:t>Anon note</w:t></w:r></w:p></w:comment>'
        )
        docx = self.build_docx(
            "legacy-defaults.docx", body, comments_xml=comments
        )
        payload = self.inspect_payload(docx)

        self.assertEqual(
            payload["comments"],
            [{"author": "Unknown", "date": "Unknown", "text": "Anon note"}],
        )
        self.assertEqual(
            payload["revisions"],
            {
                "insertions": [
                    {"author": "Unknown", "date": "Unknown", "text": "New"}
                ],
                "deletions": [],
            },
        )

        review = self.review_of(payload)
        revisions = self.list_of(review, "revisions")
        self.assertEqual(len(revisions), 1)
        self.assertEqual(revisions[0]["id"], "50")
        self.assertIsNone(revisions[0]["author"])
        self.assertIsNone(revisions[0]["date"])
        self.assert_missing_metadata(revisions[0], ["author", "date"])
        self.assertEqual(
            revisions[0]["ranges"],
            {"original": [4, 4], "proposed": [4, 7], "markup": [4, 7]},
        )

        comment = self.comment_by_id(self.list_of(review, "comments"), "1")
        self.assertIsNone(comment["author"])
        self.assertIsNone(comment["date"])
        self.assert_missing_metadata(comment, ["author", "date"])
        self.assertIsNone(comment["resolved"])

    def test_extract_preserves_legacy_v2_and_v1_payloads(self) -> None:
        docx = self.build_docx(
            "mixed.docx",
            MIXED_BODY,
            comments_xml=MIXED_COMMENTS,
            comments_extended_xml=MIXED_COMMENTS_EXTENDED,
        )
        v2, v1 = self.extract_payloads(docx)

        self.assertEqual(
            set(v2), {"schema", "source", "stats", "document", "review"}
        )
        self.assertEqual(v2["schema"], "docx-structure.v2")
        self.assertEqual(
            v2["stats"],
            {
                "sections": 0,
                "content": {"paragraph": 2, "list_item": 0, "table": 0},
                "headers": 0,
                "footers": 0,
                "pre_heading_nodes": 2,
            },
        )
        self.assertEqual(
            v2["document"]["pre_heading"],
            [
                {
                    "type": "paragraph",
                    "block_id": "p_1",
                    "text": "Intro",
                    "style": "Normal",
                },
                {
                    "type": "paragraph",
                    "block_id": "p_2",
                    "text": "ADG",
                    "style": "Normal",
                },
            ],
        )

        self.assertEqual(
            set(v1), {"schema", "source", "stats", "blocks", "chunks", "review"}
        )
        self.assertEqual(v1["schema"], "docx-llm-chunks.v1")
        self.assertEqual(
            v1["stats"],
            {
                "blocks": 2,
                "chunks": 1,
                "by_type": {
                    "heading": 0,
                    "paragraph": 2,
                    "table": 0,
                    "header": 0,
                    "footer": 0,
                },
            },
        )
        self.assertEqual(
            v1["blocks"],
            [
                {
                    "block_id": "p_1",
                    "type": "paragraph",
                    "section_path": "ROOT",
                    "style": "Normal",
                    "row_range": None,
                    "text": "Intro",
                },
                {
                    "block_id": "p_2",
                    "type": "paragraph",
                    "section_path": "ROOT",
                    "style": "Normal",
                    "row_range": None,
                    "text": "ADG",
                },
            ],
        )
        self.assertEqual(
            v1["chunks"],
            [
                {
                    "chunk_id": "c_0001",
                    "block_ids": ["p_1", "p_2"],
                    "section_paths": ["ROOT"],
                    "block_types": ["paragraph"],
                    "text": "[p_1] Intro\n\n[p_2] ADG",
                    "char_count": 22,
                }
            ],
        )
        self.assertEqual(self.review_of(v2)["schema"], REVIEW_SCHEMA)
        self.assertEqual(self.review_of(v1)["schema"], REVIEW_SCHEMA)

    def test_input_document_is_unchanged(self) -> None:
        docx = self.build_docx(
            "mixed.docx",
            MIXED_BODY,
            comments_xml=MIXED_COMMENTS,
            comments_extended_xml=MIXED_COMMENTS_EXTENDED,
        )
        before = hashlib.sha256(docx.read_bytes()).hexdigest()

        review = self.review_of(self.inspect_payload(docx))
        self.assertEqual(review["schema"], REVIEW_SCHEMA)
        self.extract_payloads(docx)

        after = hashlib.sha256(docx.read_bytes()).hexdigest()
        self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main(verbosity=2)
