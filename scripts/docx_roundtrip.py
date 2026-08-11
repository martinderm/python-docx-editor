#!/usr/bin/env python3
"""Config-driven, guarded Markdown <-> DOCX review roundtrip."""

from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SCRIPT_DIR = Path(__file__).resolve().parent
REVISION_RE = re.compile(r"^R(?P<number>\d{4})$")
CONFIG_SCHEMA = "docx-roundtrip-config.v1"
STATE_SCHEMA = "docx-roundtrip-state.v1"
MANIFEST_SCHEMA = "docx-roundtrip-manifest.v1"
REPORT_SCHEMA = "docx-roundtrip-import-report.v1"


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary, path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_docx(path: Path) -> None:
    if not path.exists() or path.suffix.lower() != ".docx":
        raise RuntimeError(f"DOCX fehlt oder hat eine falsche Endung: {path}")
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        if broken:
            raise RuntimeError(f"Beschädigter DOCX-ZIP-Eintrag: {broken}")
    Document(str(path))


def require_mapping(value: Any, name: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise RuntimeError(f"Konfigurationsfeld '{name}' muss ein JSON-Objekt sein")
    return value


def require_string(mapping: dict[str, Any], key: str, context: str = "config") -> str:
    value = mapping.get(key)
    if not isinstance(value, str) or not value.strip():
        raise RuntimeError(f"{context}.{key} muss eine nichtleere Zeichenkette sein")
    return value


def validate_config_data(config: dict[str, Any]) -> None:
    if config.get("schema") != CONFIG_SCHEMA:
        raise RuntimeError(f"config.schema muss '{CONFIG_SCHEMA}' sein")
    document_id = require_string(config, "document_id")
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,62}[a-z0-9]", document_id):
        raise RuntimeError("config.document_id muss ein stabiler Kleinbuchstaben-Slug sein")
    for key in ("canonical_markdown", "template", "working_directory"):
        require_string(config, key)

    cloud = require_mapping(config.get("cloud"), "cloud")
    for key in ("directory", "review_pattern", "archive_directory"):
        require_string(cloud, key, "cloud")
    if cloud["review_pattern"].count("{revision}") != 1:
        raise RuntimeError("cloud.review_pattern muss genau einen {revision}-Platzhalter enthalten")

    conversion = require_mapping(config.get("conversion", {}), "conversion")
    orientation = conversion.get("table_orientation", "portrait")
    if orientation not in {"portrait", "landscape"}:
        raise RuntimeError("conversion.table_orientation muss portrait oder landscape sein")
    prefix = conversion.get("bookmark_prefix", document_id.replace("-", "_"))
    if not isinstance(prefix, str) or not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]{0,15}", prefix):
        raise RuntimeError("conversion.bookmark_prefix muss ein kurzer OOXML-tauglicher Name sein")

    rules = conversion.get("merge_sections_into_landscape", [])
    if not isinstance(rules, list):
        raise RuntimeError("conversion.merge_sections_into_landscape muss eine Liste sein")
    for index, rule in enumerate(rules):
        rule = require_mapping(rule, f"conversion.merge_sections_into_landscape[{index}]")
        require_string(rule, "heading", f"conversion.merge_sections_into_landscape[{index}]")
        keep = rule.get("keep_following_paragraphs", 1)
        if not isinstance(keep, int) or keep < 0 or keep > 20:
            raise RuntimeError("keep_following_paragraphs muss zwischen 0 und 20 liegen")

    policy = require_mapping(config.get("import_policy", {}), "import_policy")
    for key in (
        "automatic_overwrite",
        "require_candidate_review",
        "refuse_if_canonical_changed_since_export",
        "high_risk_requires_explicit_override",
    ):
        if key in policy and not isinstance(policy[key], bool):
            raise RuntimeError(f"import_policy.{key} muss true oder false sein")
    if policy.get("automatic_overwrite", False):
        raise RuntimeError("automatic_overwrite=true wird aus Sicherheitsgründen nicht unterstützt")
    if not policy.get("require_candidate_review", True):
        raise RuntimeError(
            "require_candidate_review=false wird aus Sicherheitsgründen nicht unterstützt"
        )


def next_revision(revision: str) -> str:
    match = REVISION_RE.fullmatch(revision)
    if match is None:
        raise RuntimeError("Revision muss dem Muster R0001 entsprechen")
    number = int(match.group("number"))
    if number >= 9999:
        raise RuntimeError("Auf R9999 kann keine weitere vierstellige Revision folgen")
    return f"R{number + 1:04d}"


def markdown_headings(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig")
    if text.startswith("---\n"):
        _, _, text = text.partition("\n---\n")
    headings: list[dict[str, Any]] = []
    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if not match:
            continue
        title = re.sub(r"\s+#+$", "", match.group(2)).strip()
        headings.append({"level": len(match.group(1)), "title": title})
    return headings


def bookmark_name(prefix: str, level: int, title: str, occurrence: int) -> str:
    ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", ascii_title).strip("_").lower() or "section"
    token = hashlib.sha256(f"{level}|{title}|{occurrence}".encode("utf-8")).hexdigest()[:8]
    available = max(1, 40 - len(prefix) - len(token) - 2)
    return f"{prefix}_{slug[:available]}_{token}"[:40]


def inject_heading_bookmarks(
    docx_path: Path,
    headings: list[dict[str, Any]],
    prefix: str,
) -> list[dict[str, Any]]:
    doc = Document(str(docx_path))
    doc_headings = []
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        if style == "Title" and paragraph.text.strip():
            doc_headings.append((paragraph, 1, paragraph.text.strip(), style))
            continue
        match = re.match(r"(?:Heading|Überschrift)\s+(\d+)", style)
        if match and paragraph.text.strip():
            doc_headings.append(
                (paragraph, int(match.group(1)) + 1, paragraph.text.strip(), style)
            )

    if len(doc_headings) != len(headings):
        raise RuntimeError(
            f"Überschriftenzahl weicht ab: Markdown={len(headings)}, DOCX={len(doc_headings)}"
        )

    used: dict[tuple[int, str], int] = {}
    bookmark_map: list[dict[str, Any]] = []
    for index, (expected, actual) in enumerate(zip(headings, doc_headings), start=1):
        paragraph, level, title, style = actual
        if (expected["level"], expected["title"]) != (level, title):
            raise RuntimeError(
                f"Überschrift {index} weicht ab: Markdown={expected!r}, "
                f"DOCX={{'level': {level}, 'title': {title!r}}}"
            )
        key = (level, title)
        used[key] = used.get(key, 0) + 1
        name = bookmark_name(prefix, level, title, used[key])
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(1000 + index))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(1000 + index))
        paragraph._p.insert(0, start)  # pylint: disable=protected-access
        paragraph._p.append(end)  # pylint: disable=protected-access
        bookmark_map.append(
            {
                "index": index,
                "markdown_level": level,
                "word_style": style,
                "title": title,
                "bookmark": name,
            }
        )
    doc.save(str(docx_path))
    return bookmark_map


def paragraph_has_section_break(element) -> bool:
    if element is None or element.tag != qn("w:p"):
        return False
    properties = element.find(qn("w:pPr"))
    return properties is not None and properties.find(qn("w:sectPr")) is not None


def merge_section_into_landscape(
    doc: Document,
    heading: str,
    keep_following_paragraphs: int,
) -> None:
    target = next((p for p in doc.paragraphs if p.text.strip() == heading), None)
    if target is None:
        raise RuntimeError(f"Konfigurierte Layoutüberschrift wurde nicht gefunden: {heading}")

    previous = target._p.getprevious()  # pylint: disable=protected-access
    while previous is not None and not paragraph_has_section_break(previous):
        previous = previous.getprevious()
    following = target._p.getnext()  # pylint: disable=protected-access
    while following is not None and not paragraph_has_section_break(following):
        following = following.getnext()
    if previous is None or following is None:
        raise RuntimeError(f"Abschnittsgrenzen für Layoutüberschrift fehlen: {heading}")

    body = doc.element.body
    body.remove(previous)
    body.remove(following)
    target.paragraph_format.keep_with_next = True
    current = target._p.getnext()  # pylint: disable=protected-access
    kept = 0
    while current is not None and kept < keep_following_paragraphs:
        if current.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            Paragraph(current, target._parent).paragraph_format.keep_with_next = True
            kept += 1
        current = current.getnext()


def apply_configured_postprocessing(docx_path: Path, conversion: dict[str, Any]) -> None:
    rules = conversion.get("merge_sections_into_landscape", [])
    if not rules:
        return
    doc = Document(str(docx_path))
    for rule in rules:
        merge_section_into_landscape(
            doc,
            rule["heading"],
            int(rule.get("keep_following_paragraphs", 1)),
        )
    doc.save(str(docx_path))


def count_bookmarks(path: Path, prefix: str) -> int:
    doc = Document(str(path))
    xpath = f'.//w:bookmarkStart[starts-with(@w:name, "{prefix}_")]'
    return len(doc.element.body.xpath(xpath))


def split_frontmatter(text: str) -> tuple[str, str]:
    if not text.startswith("---\n"):
        return "", text
    end = text.find("\n---\n", 4)
    if end < 0:
        return "", text
    return text[: end + 5], text[end + 5 :]


def flatten_sections(data: dict[str, Any]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []

    def walk(nodes: list[dict[str, Any]], parents: list[str]) -> None:
        for node in nodes:
            path = [*parents, node["title"]]
            result.append(
                {
                    "level": node["level"],
                    "title": node["title"],
                    "path": " / ".join(path),
                    "content": node.get("content", []),
                }
            )
            walk(node.get("children", []), path)

    walk(data["document"]["sections"], [])
    return result


def section_fingerprints(data: dict[str, Any]) -> dict[str, str]:
    fingerprints: dict[str, str] = {}
    for section in flatten_sections(data):
        normalized = json.dumps(section["content"], ensure_ascii=False, sort_keys=True)
        fingerprints[section["path"]] = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    return fingerprints


class RoundtripEngine:
    def __init__(self, config_path: Path, root: Path):
        self.root = root.resolve()
        self.config_path = config_path.resolve()
        self.config = load_json(self.config_path)
        validate_config_data(self.config)
        self.document_id = self.config["document_id"]
        self.working_directory = self.path(self.config["working_directory"])
        self.state_path = self.working_directory / "state.json"
        self.lock_path = self.working_directory / "operation.lock"
        self.conversion = self.config.get("conversion", {})
        self.policy = self.config.get("import_policy", {})
        self.bookmark_prefix = self.conversion.get(
            "bookmark_prefix", self.document_id.replace("-", "_")
        )

    def path(self, value: str) -> Path:
        path = Path(value)
        return path if path.is_absolute() else self.root / path

    def relative_path(self, path: Path) -> str:
        try:
            return str(path.relative_to(self.root)).replace("\\", "/")
        except ValueError:
            return str(path)

    def tool(self, key: str, default_name: str) -> Path:
        configured = self.conversion.get(key)
        return self.path(configured) if configured else SCRIPT_DIR / default_name

    def run_python(self, script: Path, *arguments: str) -> None:
        result = subprocess.run(
            [sys.executable, str(script), *arguments],
            cwd=self.root,
            check=False,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
        )
        if result.returncode:
            raise RuntimeError(
                f"Hilfsskript fehlgeschlagen ({script.name}):\n{result.stdout}\n{result.stderr}"
            )

    @contextmanager
    def operation_lock(self, operation: str) -> Iterator[None]:
        self.working_directory.mkdir(parents=True, exist_ok=True)
        payload = {
            "schema": "docx-roundtrip-lock.v1",
            "document_id": self.document_id,
            "operation": operation,
            "created_at": now_iso(),
            "pid": os.getpid(),
        }
        try:
            descriptor = os.open(self.lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError as exc:
            existing = self.lock_path.read_text(encoding="utf-8", errors="replace")
            raise RuntimeError(
                f"Roundtrip ist gesperrt: {self.lock_path}\n{existing}"
            ) from exc
        try:
            with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
                json.dump(payload, handle, ensure_ascii=False, indent=2)
                handle.write("\n")
            yield
        finally:
            if self.lock_path.exists():
                self.lock_path.unlink()

    def revision_paths(self, revision: str, *, base: Path | None = None) -> dict[str, Path]:
        revision_dir = (base or (self.working_directory / "revisions")) / revision
        cloud = self.config["cloud"]
        cloud_dir = self.path(cloud["directory"])
        cloud_name = cloud["review_pattern"].format(revision=revision)
        return {
            "revision_dir": revision_dir,
            "base_md": revision_dir / "base.md",
            "export_docx": revision_dir / "export.docx",
            "base_extracted": revision_dir / "base-extracted.md",
            "base_structure": revision_dir / "base-structure.json",
            "manifest": revision_dir / "manifest.json",
            "cloud_docx": cloud_dir / cloud_name,
        }

    def state(self) -> dict[str, Any]:
        if not self.state_path.exists():
            raise RuntimeError(f"Roundtrip-Status fehlt: {self.state_path}")
        state = load_json(self.state_path)
        if state.get("schema") != STATE_SCHEMA:
            raise RuntimeError(f"state.schema muss '{STATE_SCHEMA}' sein")
        if state.get("document_id") != self.document_id:
            raise RuntimeError("Statusdatei gehört zu einem anderen document_id")
        revision = state.get("current_revision")
        if not isinstance(revision, str) or REVISION_RE.fullmatch(revision) is None:
            raise RuntimeError("Statusdatei enthält keine gültige current_revision")
        expected_cloud = self.revision_paths(revision)["cloud_docx"].resolve()
        actual_cloud_value = state.get("cloud_docx")
        if not isinstance(actual_cloud_value, str):
            raise RuntimeError("Statusdatei enthält keinen gültigen cloud_docx-Pfad")
        actual_cloud = self.path(actual_cloud_value).resolve()
        if actual_cloud != expected_cloud:
            raise RuntimeError(
                "Statusdatei verweist nicht auf den für current_revision konfigurierten Cloudpfad"
            )
        return state

    def review_filename_regex(self) -> re.Pattern[str]:
        pattern = self.config["cloud"]["review_pattern"]
        expression = re.escape(pattern).replace(
            re.escape("{revision}"),
            r"(?P<revision>R\d{4})",
        )
        return re.compile(f"^{expression}$", re.IGNORECASE)

    def cloud_inventory(self, state: dict[str, Any]) -> dict[str, Any]:
        cloud_dir = self.path(self.config["cloud"]["directory"])
        current_revision = state["current_revision"]
        expected = self.path(state["cloud_docx"])
        recognized: list[dict[str, str]] = []
        unrecognized: list[str] = []
        if cloud_dir.exists():
            filename_re = self.review_filename_regex()
            for path in sorted(cloud_dir.iterdir(), key=lambda item: item.name.casefold()):
                if not path.is_file() or path.suffix.lower() != ".docx" or path.name.startswith("."):
                    continue
                match = filename_re.fullmatch(path.name)
                if match is None:
                    unrecognized.append(self.relative_path(path))
                    continue
                recognized.append(
                    {
                        "revision": match.group("revision").upper(),
                        "path": self.relative_path(path),
                    }
                )
        other_revisions = [item for item in recognized if item["revision"] != current_revision]
        warnings: list[str] = []
        if other_revisions:
            warnings.append(
                "Unerwartete zusätzliche Reviewrevision(en): "
                + ", ".join(item["path"] for item in other_revisions)
            )
        if unrecognized:
            warnings.append(
                "Nicht erkannte DOCX-Datei(en) im Reviewordner: "
                + ", ".join(unrecognized)
            )
        return {
            "expected": self.relative_path(expected),
            "expected_exists": expected.exists(),
            "recognized": recognized,
            "other_revisions": other_revisions,
            "unrecognized_docx": unrecognized,
            "warnings": warnings,
        }

    def require_unambiguous_review(self, state: dict[str, Any]) -> None:
        inventory = self.cloud_inventory(state)
        if not inventory["expected_exists"]:
            raise RuntimeError("Die aktuelle Reviewdatei fehlt; Dateinamen und Cloud-Sync klären")
        if inventory["other_revisions"] or inventory["unrecognized_docx"]:
            raise RuntimeError(
                "Der Reviewordner ist nicht eindeutig: " + " | ".join(inventory["warnings"])
            )

    @staticmethod
    def classify_hash(state: dict[str, Any], current_hash: str) -> str:
        imported_hash = state.get("last_imported_docx_sha256")
        if imported_hash and current_hash == imported_hash:
            return "unchanged_since_import"
        if current_hash == state["docx_sha256_at_export"]:
            return "unchanged"
        if (
            state.get("status") == "candidate_prepared"
            and current_hash == state.get("last_seen_docx_sha256")
        ):
            return "candidate_prepared"
        return "modified"

    def extract_docx(self, docx: Path, markdown_out: Path, structure_out: Path) -> None:
        self.run_python(
            self.tool("docx_to_markdown", "docx_to_markdown.py"),
            "--in",
            str(docx),
            "--out",
            str(markdown_out),
        )
        self.run_python(
            self.tool("extract_structure", "extract_docx_for_llm.py"),
            "--in",
            str(docx),
            "--out",
            str(structure_out),
        )

    def export_revision(self, revision: str, *, replace: bool = False) -> dict[str, Any]:
        if REVISION_RE.fullmatch(revision) is None:
            raise RuntimeError("Revision muss dem Muster R0001 entsprechen")
        canonical = self.path(self.config["canonical_markdown"])
        template = self.path(self.config["template"])
        paths = self.revision_paths(revision)
        if paths["revision_dir"].exists() and not replace:
            raise RuntimeError(f"Revision existiert bereits: {paths['revision_dir']}")
        if paths["cloud_docx"].exists() and not replace:
            raise RuntimeError(f"Cloudrevision existiert bereits: {paths['cloud_docx']}")
        if replace:
            raise RuntimeError("replace wird in der generischen Engine nicht unterstützt")

        staging_root = self.working_directory / "staging"
        staging_revision = staging_root / revision
        if staging_revision.exists():
            raise RuntimeError(f"Unbereinigtes Staging-Verzeichnis vorhanden: {staging_revision}")
        staging_paths = self.revision_paths(revision, base=staging_root)
        staging_revision.mkdir(parents=True, exist_ok=False)
        cloud_created = False
        final_created = False
        try:
            shutil.copy2(canonical, staging_paths["base_md"])
            with tempfile.TemporaryDirectory(prefix="docx-roundtrip-") as temp_dir:
                generated = Path(temp_dir) / "generated.docx"
                self.run_python(
                    self.tool("markdown_to_docx", "markdown_to_template_docx.py"),
                    "--in",
                    str(canonical),
                    "--out",
                    str(generated),
                    "--template",
                    str(template),
                    "--table-orientation",
                    self.conversion.get("table_orientation", "portrait"),
                )
                bookmarks = inject_heading_bookmarks(
                    generated,
                    markdown_headings(canonical),
                    self.bookmark_prefix,
                )
                apply_configured_postprocessing(generated, self.conversion)
                validate_docx(generated)
                shutil.copy2(generated, staging_paths["export_docx"])

            self.extract_docx(
                staging_paths["export_docx"],
                staging_paths["base_extracted"],
                staging_paths["base_structure"],
            )
            exported_at = now_iso()
            canonical_hash = sha256(canonical)
            export_hash = sha256(staging_paths["export_docx"])
            manifest = {
                "schema": MANIFEST_SCHEMA,
                "document_id": self.document_id,
                "revision": revision,
                "exported_at": exported_at,
                "paths": {
                    "canonical_markdown": self.relative_path(canonical),
                    "cloud_docx": self.relative_path(paths["cloud_docx"]),
                    "template": self.relative_path(template),
                },
                "hashes": {
                    "canonical_markdown": canonical_hash,
                    "base_markdown": sha256(staging_paths["base_md"]),
                    "export_docx": export_hash,
                    "cloud_docx_at_export": export_hash,
                },
                "headings": bookmarks,
            }
            write_json(staging_paths["manifest"], manifest)

            paths["cloud_docx"].parent.mkdir(parents=True, exist_ok=True)
            self.path(self.config["cloud"]["archive_directory"]).mkdir(
                parents=True, exist_ok=True
            )
            temporary_cloud = paths["cloud_docx"].with_name(
                f".{paths['cloud_docx'].name}.tmp"
            )
            shutil.copy2(staging_paths["export_docx"], temporary_cloud)
            os.replace(temporary_cloud, paths["cloud_docx"])
            cloud_created = True

            paths["revision_dir"].parent.mkdir(parents=True, exist_ok=True)
            os.replace(staging_revision, paths["revision_dir"])
            final_created = True
            state = {
                "schema": STATE_SCHEMA,
                "document_id": self.document_id,
                "current_revision": revision,
                "cloud_docx": manifest["paths"]["cloud_docx"],
                "exported_at": exported_at,
                "canonical_sha256_at_export": canonical_hash,
                "docx_sha256_at_export": export_hash,
                "last_seen_docx_sha256": export_hash,
                "last_imported_docx_sha256": None,
                "status": "awaiting_review",
            }
            write_json(self.state_path, state)
            return {
                "status": "exported",
                "revision": revision,
                "cloud_docx": manifest["paths"]["cloud_docx"],
                "bookmarks": len(bookmarks),
            }
        except Exception:
            if cloud_created and paths["cloud_docx"].exists():
                paths["cloud_docx"].unlink()
            if final_created and paths["revision_dir"].exists():
                shutil.rmtree(paths["revision_dir"])
            if staging_revision.exists():
                shutil.rmtree(staging_revision)
            raise

    def status(self) -> dict[str, Any]:
        state = self.state()
        cloud_docx = self.path(state["cloud_docx"])
        inventory = self.cloud_inventory(state)
        if not cloud_docx.exists():
            return {
                "status": "missing_current_revision",
                "revision": state["current_revision"],
                "cloud_docx": state["cloud_docx"],
                "inventory": inventory,
                "warnings": inventory["warnings"],
            }
        validate_docx(cloud_docx)
        current_hash = sha256(cloud_docx)
        return {
            "status": self.classify_hash(state, current_hash),
            "revision": state["current_revision"],
            "cloud_docx": state["cloud_docx"],
            "sha256": current_hash,
            "bookmarks": count_bookmarks(cloud_docx, self.bookmark_prefix),
            "inventory": inventory,
            "warnings": inventory["warnings"],
        }

    def prepare_import(self) -> dict[str, Any]:
        state = self.state()
        self.require_unambiguous_review(state)
        revision = state["current_revision"]
        paths = self.revision_paths(revision)
        cloud_docx = self.path(state["cloud_docx"])
        validate_docx(cloud_docx)
        incoming_hash = sha256(cloud_docx)
        if incoming_hash == state.get("last_imported_docx_sha256"):
            return {"status": "already_imported"}
        if incoming_hash == state["docx_sha256_at_export"]:
            return {"status": "unchanged"}
        report_path = paths["revision_dir"] / "import-report.json"
        if (
            state.get("status") == "candidate_prepared"
            and incoming_hash == state.get("last_seen_docx_sha256")
            and report_path.exists()
        ):
            return {"status": "candidate_already_prepared", "revision": revision}

        incoming_docx = paths["revision_dir"] / "incoming.docx"
        incoming_md = paths["revision_dir"] / "incoming-extracted.md"
        incoming_structure = paths["revision_dir"] / "incoming-structure.json"
        candidate = paths["revision_dir"] / "candidate.md"
        diff_path = paths["revision_dir"] / "changes.diff"
        shutil.copy2(cloud_docx, incoming_docx)
        self.extract_docx(incoming_docx, incoming_md, incoming_structure)

        canonical = self.path(self.config["canonical_markdown"])
        canonical_text = canonical.read_text(encoding="utf-8-sig")
        preserve_frontmatter = self.conversion.get("preserve_frontmatter", True)
        frontmatter, _ = split_frontmatter(canonical_text) if preserve_frontmatter else ("", "")
        candidate.write_text(
            frontmatter + incoming_md.read_text(encoding="utf-8-sig").rstrip() + "\n",
            encoding="utf-8",
        )
        base_text = paths["base_extracted"].read_text(encoding="utf-8-sig").splitlines(
            keepends=True
        )
        incoming_text = incoming_md.read_text(encoding="utf-8-sig").splitlines(keepends=True)
        diff = difflib.unified_diff(
            base_text,
            incoming_text,
            fromfile="base-extracted.md",
            tofile="incoming-extracted.md",
        )
        diff_path.write_text("".join(diff), encoding="utf-8")

        base_structure = load_json(paths["base_structure"])
        edited_structure = load_json(incoming_structure)
        base_sections = flatten_sections(base_structure)
        edited_sections = flatten_sections(edited_structure)
        base_sequence = [(s["level"], s["title"]) for s in base_sections]
        edited_sequence = [(s["level"], s["title"]) for s in edited_sections]
        base_fp = section_fingerprints(base_structure)
        edited_fp = section_fingerprints(edited_structure)
        changed_sections = sorted(
            path
            for path in set(base_fp) | set(edited_fp)
            if base_fp.get(path) != edited_fp.get(path)
        )
        base_stats = base_structure["stats"]
        edited_stats = edited_structure["stats"]
        bookmark_count = count_bookmarks(incoming_docx, self.bookmark_prefix)
        expected_bookmarks = len(load_json(paths["manifest"])["headings"])
        if (
            base_sequence != edited_sequence
            or bookmark_count < expected_bookmarks
            or base_stats["content"]["table"] != edited_stats["content"]["table"]
        ):
            risk = "high"
        elif base_stats["content"] != edited_stats["content"]:
            risk = "medium"
        else:
            risk = "low"
        report = {
            "schema": REPORT_SCHEMA,
            "document_id": self.document_id,
            "revision": revision,
            "prepared_at": now_iso(),
            "incoming_docx_sha256": incoming_hash,
            "risk": risk,
            "structure": {
                "heading_sequence_unchanged": base_sequence == edited_sequence,
                "bookmarks_expected": expected_bookmarks,
                "bookmarks_found": bookmark_count,
                "base_stats": base_stats,
                "incoming_stats": edited_stats,
                "changed_sections": changed_sections,
            },
            "artifacts": {
                "candidate": "candidate.md",
                "diff": "changes.diff",
                "incoming_docx": "incoming.docx",
                "incoming_markdown": "incoming-extracted.md",
                "incoming_structure": "incoming-structure.json",
            },
            "decision": "candidate_review_required",
        }
        write_json(report_path, report)
        state["last_seen_docx_sha256"] = incoming_hash
        state["status"] = "candidate_prepared"
        state["candidate_risk"] = risk
        write_json(self.state_path, state)
        return {
            "status": "candidate_prepared",
            "revision": revision,
            "risk": risk,
            "changed_sections": changed_sections,
            "candidate": self.relative_path(candidate),
        }

    def apply_import(self, *, allow_high_risk: bool) -> dict[str, Any]:
        state = self.state()
        self.require_unambiguous_review(state)
        revision = state["current_revision"]
        paths = self.revision_paths(revision)
        report = load_json(paths["revision_dir"] / "import-report.json")
        candidate = paths["revision_dir"] / "candidate.md"
        canonical = self.path(self.config["canonical_markdown"])
        cloud_docx = self.path(state["cloud_docx"])
        validate_docx(cloud_docx)
        if sha256(cloud_docx) != report["incoming_docx_sha256"]:
            raise RuntimeError(
                "Word-Datei wurde nach prepare-import verändert; Kandidaten neu vorbereiten"
            )
        if (
            self.policy.get("refuse_if_canonical_changed_since_export", True)
            and sha256(canonical) != state["canonical_sha256_at_export"]
        ):
            raise RuntimeError("Kanonisches Markdown wurde seit dem Export verändert")
        if (
            report["risk"] == "high"
            and self.policy.get("high_risk_requires_explicit_override", True)
            and not allow_high_risk
        ):
            raise RuntimeError("Strukturänderung mit Risiko high; --allow-high-risk erforderlich")
        backup = paths["revision_dir"] / "pre-import.md"
        shutil.copy2(canonical, backup)
        temporary = canonical.with_name(f".{canonical.name}.tmp")
        shutil.copy2(candidate, temporary)
        os.replace(temporary, canonical)
        state["last_imported_docx_sha256"] = report["incoming_docx_sha256"]
        state["last_seen_docx_sha256"] = report["incoming_docx_sha256"]
        state["status"] = "imported"
        state["imported_at"] = now_iso()
        state["canonical_sha256_after_import"] = sha256(canonical)
        write_json(self.state_path, state)
        return {
            "status": "imported",
            "revision": revision,
            "canonical": self.relative_path(canonical),
            "backup": self.relative_path(backup),
        }

    def advance_revision(self) -> dict[str, Any]:
        state = self.state()
        self.require_unambiguous_review(state)
        current_revision = state["current_revision"]
        current_docx = self.path(state["cloud_docx"])
        validate_docx(current_docx)
        current_hash = sha256(current_docx)
        cloud_status = self.classify_hash(state, current_hash)
        if cloud_status not in {"unchanged", "unchanged_since_import"}:
            raise RuntimeError(
                f"Reviewdatei hat Status '{cloud_status}'; Änderungen zuerst abschließen"
            )
        canonical = self.path(self.config["canonical_markdown"])
        expected_canonical_hash = (
            state.get("canonical_sha256_after_import")
            if cloud_status == "unchanged_since_import"
            else state.get("canonical_sha256_at_export")
        )
        if not expected_canonical_hash or sha256(canonical) != expected_canonical_hash:
            raise RuntimeError("Kanonisches Markdown wurde seit dem letzten Abschluss verändert")

        revision = next_revision(current_revision)
        next_paths = self.revision_paths(revision)
        if next_paths["revision_dir"].exists() or next_paths["cloud_docx"].exists():
            raise RuntimeError(f"Zielrevision {revision} existiert bereits")
        archive_dir = self.path(self.config["cloud"]["archive_directory"])
        archive_dir.mkdir(parents=True, exist_ok=True)
        archive_docx = archive_dir / current_docx.name
        if archive_docx.exists():
            raise RuntimeError(f"Archivziel existiert bereits: {archive_docx}")

        os.replace(current_docx, archive_docx)
        try:
            result = self.export_revision(revision)
        except Exception:
            if not current_docx.exists() and archive_docx.exists():
                os.replace(archive_docx, current_docx)
            raise
        new_state = self.state()
        new_state["previous_revision"] = current_revision
        new_state["previous_archive_docx"] = self.relative_path(archive_docx)
        new_state["advanced_at"] = now_iso()
        write_json(self.state_path, new_state)
        result.update(
            {
                "status": "revision_advanced",
                "previous_revision": current_revision,
                "archive_docx": self.relative_path(archive_docx),
            }
        )
        return result


def print_result(result: dict[str, Any]) -> None:
    print(json.dumps(result, ensure_ascii=False, indent=2))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", required=True, help="Roundtrip JSON configuration")
    parser.add_argument(
        "--root",
        default=".",
        help="Base directory for relative paths in the config (default: current directory)",
    )
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("validate", help="Konfiguration und referenzierte Pfade prüfen")
    export = sub.add_parser("export", help="Kanonisches Markdown als Review-DOCX exportieren")
    export.add_argument("--revision", required=True)
    export.add_argument("--confirm", action="store_true")
    sub.add_parser("status", help="Aktuelle Cloudrevision und Dateiinventur prüfen")
    sub.add_parser("prepare-import", help="Geänderte Word-Fassung als Kandidat vorbereiten")
    apply_cmd = sub.add_parser("apply-import", help="Geprüften Kandidaten übernehmen")
    apply_cmd.add_argument("--confirm", action="store_true")
    apply_cmd.add_argument("--allow-high-risk", action="store_true")
    advance = sub.add_parser("advance-revision", help="Archivieren und Folgerevision erzeugen")
    advance.add_argument("--confirm", action="store_true")
    return parser


def validate_referenced_paths(engine: RoundtripEngine) -> dict[str, Any]:
    canonical = engine.path(engine.config["canonical_markdown"])
    template = engine.path(engine.config["template"])
    errors = []
    if not canonical.exists():
        errors.append(f"Kanonisches Markdown fehlt: {canonical}")
    if not template.exists():
        errors.append(f"Word-Vorlage fehlt: {template}")
    for key, default in (
        ("markdown_to_docx", "markdown_to_template_docx.py"),
        ("docx_to_markdown", "docx_to_markdown.py"),
        ("extract_structure", "extract_docx_for_llm.py"),
    ):
        tool = engine.tool(key, default)
        if not tool.exists():
            errors.append(f"Hilfsskript fehlt: {tool}")
    if errors:
        raise RuntimeError("\n".join(errors))
    return {
        "status": "valid",
        "schema": engine.config["schema"],
        "document_id": engine.document_id,
        "canonical_markdown": engine.relative_path(canonical),
        "working_directory": engine.relative_path(engine.working_directory),
    }


def main() -> int:
    args = build_parser().parse_args()
    try:
        engine = RoundtripEngine(Path(args.config), Path(args.root))
        if args.command == "validate":
            print_result(validate_referenced_paths(engine))
            return 0
        if args.command == "status":
            result = engine.status()
            print_result(result)
            return 2 if result["status"] == "missing_current_revision" else 0
        if args.command == "export":
            if not args.confirm:
                raise RuntimeError("Export in den Cloudordner nur mit --confirm")
            with engine.operation_lock("export"):
                if engine.state_path.exists():
                    state = engine.state()
                    if args.revision != state.get("current_revision"):
                        raise RuntimeError(
                            "Aktive Revision vorhanden; advance-revision --confirm verwenden"
                        )
                print_result(engine.export_revision(args.revision))
            return 0
        if args.command == "prepare-import":
            with engine.operation_lock("prepare-import"):
                print_result(engine.prepare_import())
            return 0
        if args.command == "apply-import":
            if not args.confirm:
                raise RuntimeError("Übernahme nur mit --confirm")
            with engine.operation_lock("apply-import"):
                print_result(engine.apply_import(allow_high_risk=args.allow_high_risk))
            return 0
        if args.command == "advance-revision":
            if not args.confirm:
                raise RuntimeError("Revisionswechsel nur mit --confirm")
            with engine.operation_lock("advance-revision"):
                print_result(engine.advance_revision())
            return 0
        raise RuntimeError(f"Unbekannter Befehl: {args.command}")
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
